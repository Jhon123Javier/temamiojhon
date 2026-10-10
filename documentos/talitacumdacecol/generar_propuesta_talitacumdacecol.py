#!/usr/bin/env python3
"""Propuesta comercial Talitacumdacecol con membrete de The Company Dev."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor, Twips

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUT = ROOT / "Propuesta The Company Dev - Talitacumdacecol - Sitio Web y Redes.docx"

NAVY = RGBColor(0x00, 0x30, 0x60)
BLACK = RGBColor(0x1B, 0x33, 0x25)
GREEN = "02AF5D"
DARK = "1B3325"


def set_run_font(run, size=13, bold=False, color=None, name="Poppins"):
    run.bold = bold
    run.italic = False
    run.font.size = Pt(size)
    run.font.name = name
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), name)


def set_paragraph_format(paragraph, before=0, after=8, align="justify", keep_next=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = 1.08
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if align == "justify":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "center":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pPr = paragraph._p.get_or_add_pPr()
    if keep_next:
        if pPr.find(qn("w:keepNext")) is None:
            pPr.append(OxmlElement("w:keepNext"))
    if pPr.find(qn("w:widowControl")) is None:
        pPr.append(OxmlElement("w:widowControl"))


def add_runs(paragraph, parts, size=13, color=None):
    for text, bold in parts:
        run = paragraph.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color if color else BLACK)


def add_body(doc, parts, size=13, before=0, after=8, align="justify", keep_next=False):
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, before=before, after=after, align=align, keep_next=keep_next)
    add_runs(paragraph, parts, size=size)
    return paragraph


def add_heading(doc, text):
    return add_body(
        doc,
        [(text, True)],
        size=13,
        before=12,
        after=4,
        align="left",
        keep_next=True,
    )


def add_subheading(doc, text):
    return add_body(
        doc,
        [(text, True)],
        size=13,
        before=8,
        after=2,
        align="left",
        keep_next=True,
    )


def add_bullet(doc, parts, size=13):
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, before=1, after=2, align="left")
    fmt = paragraph.paragraph_format
    fmt.left_indent = Cm(0.75)
    fmt.first_line_indent = Cm(-0.4)
    add_runs(paragraph, [("•  ", False)] + parts, size=size)
    return paragraph


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_borders(cell, color="1B3325", sz="4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), sz)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        tcBorders.append(element)
    tcPr.append(tcBorders)


def set_cell_margins(cell, top=40, bottom=40, left=80, right=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def prevent_row_split(row, keep_with_next=False):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    if trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))
    if keep_with_next:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                pPr = paragraph._p.get_or_add_pPr()
                if pPr.find(qn("w:keepNext")) is None:
                    pPr.append(OxmlElement("w:keepNext"))


def write_cell(cell, text, bold=False, align="left", fill=None, size=12, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    set_paragraph_format(paragraph, before=0, after=0, align=align)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color or BLACK)
    set_cell_borders(cell)
    set_cell_margins(cell)
    if fill:
        shade_cell(cell, fill)
    # vertical align center
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    vAlign = OxmlElement("w:vAlign")
    vAlign.set(qn("w:val"), "center")
    tcPr.append(vAlign)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    total = sum(widths)
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    # fixed layout
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
    else:
        grid = OxmlElement("w:tblGrid")
        tblPr.addnext(grid)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        write_cell(cell, header, bold=True, align="center", fill=DARK, size=11, color=RGBColor(255, 255, 255))
        cell.width = Twips(widths[index])
    prevent_row_split(table.rows[0], keep_with_next=True)

    for r_index, row in enumerate(rows):
        for c_index, value in enumerate(row):
            align = "center" if c_index != 1 else "left"
            cell = table.rows[r_index + 1].cells[c_index]
            write_cell(cell, value, bold=False, align=align, fill="F7FBF8" if r_index % 2 else "FFFFFF", size=11)
            cell.width = Twips(widths[c_index])
        prevent_row_split(table.rows[r_index + 1], keep_with_next=r_index < len(rows) - 1)

    # space after table
    spacer = doc.add_paragraph()
    set_paragraph_format(spacer, before=0, after=2, align="left")
    return table


def add_background(header, image_path):
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_format(paragraph, before=0, after=0, align="left")
    run = paragraph.add_run()
    set_run_font(run, size=1)
    run.add_picture(image_path, width=Inches(8.5), height=Inches(11))
    drawing = run._element.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))
    extent = inline.find(qn("wp:extent"))
    cx = extent.get("cx")
    cy = extent.get("cy")
    doc_pr = inline.find(qn("wp:docPr"))
    graphic = inline.find(qn("a:graphic"))
    anchor = parse_xml(
        f'''
        <wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
                   xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
                   distT="0" distB="0" distL="0" distR="0"
                   simplePos="0" relativeHeight="0" behindDoc="1"
                   locked="0" layoutInCell="1" allowOverlap="1">
          <wp:simplePos x="0" y="0"/>
          <wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>
          <wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>
          <wp:extent cx="{cx}" cy="{cy}"/>
          <wp:effectExtent l="0" t="0" r="0" b="0"/>
          <wp:wrapNone/>
          <wp:docPr id="{doc_pr.get("id")}" name="membrete"/>
          <wp:cNvGraphicFramePr>
            <a:graphicFrameLocks xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" noChangeAspect="1"/>
          </wp:cNvGraphicFramePr>
        </wp:anchor>
        '''
    )
    anchor.append(graphic)
    drawing.replace(inline, anchor)


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.72)
    section.bottom_margin = Inches(1.08)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.15)
    section.footer_distance = Inches(0.15)
    section.different_first_page_header_footer = False
    sectPr = section._sectPr
    # one header for every page
    for tag in ("w:titlePg",):
        node = sectPr.find(qn(tag))
        if node is not None:
            sectPr.remove(node)


def add_label_line(doc, label, value):
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, before=0, after=1, align="left")
    label_run = paragraph.add_run(label)
    set_run_font(label_run, size=13, bold=True, color=NAVY)
    value_run = paragraph.add_run(value)
    set_run_font(value_run, size=13, bold=False, color=BLACK)
    # tab-like spacing via two runs; pad label
    return paragraph


def add_contacts(doc):
    contacts = [
        (ASSETS / "icon_telefono.png", "3107862438 - 3118620020"),
        (ASSETS / "icon_correo.png", "info@thecompanydev.com"),
        (ASSETS / "icon_web.png", "www.thecompanydev.com"),
    ]
    for index, (icon, text) in enumerate(contacts):
        paragraph = doc.add_paragraph()
        set_paragraph_format(paragraph, before=14 if index == 0 else 1, after=0, align="right")
        if index < len(contacts) - 1:
            pPr = paragraph._p.get_or_add_pPr()
            if pPr.find(qn("w:keepNext")) is None:
                pPr.append(OxmlElement("w:keepNext"))
        icon_run = paragraph.add_run()
        icon_run.add_picture(str(icon), width=Inches(0.2), height=Inches(0.2))
        text_run = paragraph.add_run("  " + text)
        set_run_font(text_run, size=11, bold=False, color=BLACK)


def build():
    doc = Document()
    configure_section(doc.sections[0])
    add_background(doc.sections[0].header, str(ASSETS / "membrete_fondo_carta_sin_contacto.png"))

    normal = doc.styles["Normal"]
    normal.font.name = "Poppins"
    normal.font.size = Pt(13)
    normal.font.color.rgb = BLACK
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), "Poppins")

    add_label_line(doc, "Fecha     ", "10 de octubre de 2026")
    add_label_line(doc, "Asunto   ", "Propuesta de Prestación de Servicios")
    add_label_line(doc, "Para      ", "Talitacumdacecol")

    add_body(
        doc,
        [("PROPUESTA COMERCIAL", True)],
        size=14,
        before=10,
        after=0,
        align="left",
        keep_next=True,
    )
    add_body(
        doc,
        [("SITIO WEB CORPORATIVO Y PLAN DE REDES SOCIALES", True)],
        size=13,
        before=0,
        after=8,
        align="left",
    )
    title = doc.paragraphs[-1]
    for run in title.runs:
        set_run_font(run, size=13, bold=True, color=NAVY)

    add_body(
        doc,
        [
            ("Por medio de la presente, ", False),
            ("THE COMPANY DEV", True),
            (" presenta a Talitacumdacecol la propuesta para su sitio web corporativo y para el manejo mensual de sus redes sociales. ", False),
            ("EL CLIENTE", True),
            (" puede tomar uno de los dos frentes o los dos. El trabajo queda listo para crecer: más adelante puede pedir campañas, más piezas o funciones nuevas.", False),
        ],
    )
    add_body(
        doc,
        [
            ("Sitio web corporativo. ", True),
            ("Presenta la empresa, sus servicios y sus datos de contacto. Valor: ", False),
            ("$800.000 COP", True),
            (".", False),
        ],
    )
    add_body(
        doc,
        [
            ("Plan de redes sociales. ", True),
            ("Cada mes incluye ayuda de guiones, edición y subida de ", False),
            ("14 videos", True),
            (", y edición y publicación de ", False),
            ("16 imágenes", True),
            (" en todas las redes. Valor: ", False),
            ("$1.000.000 COP al mes", True),
            (".", False),
        ],
    )
    add_body(
        doc,
        [
            ("Campañas. ", True),
            ("La consultoría y la asesoría de una campaña valen ", False),
            ("$50.000 COP", True),
            (". Realizar cada campaña vale ", False),
            ("$120.000 COP", True),
            (".", False),
        ],
    )

    add_heading(doc, "1. Objetivo del servicio")
    add_body(
        doc,
        [
            ("Dejar a Talitacumdacecol con un sitio web corporativo claro y, si lo toma, con un plan mensual que publique videos e imágenes en sus redes. ", False),
            ("EL CLIENTE", True),
            (" también puede pedir consultoría de campañas o que ", False),
            ("THE COMPANY DEV", True),
            (" realice cada campaña.", False),
        ],
    )

    add_heading(doc, "2. Alcance del servicio")
    add_subheading(doc, "2.1 Sitio web corporativo")
    add_body(
        doc,
        [
            ("Valor del proyecto: ", False),
            ("$800.000 COP", True),
            (". El calendario de publicación se fija el día en que Talitacumdacecol entregue la información completa.", False),
        ],
        after=4,
    )
    add_body(
        doc,
        [
            ("La página presenta quién es la empresa, qué ofrece y cómo contactarla. Queda administrable y se ve bien en celular y en computador.", False),
        ],
    )
    add_body(doc, [("El sitio incluye:", False)], before=2, after=2, align="left", keep_next=True)
    for item in [
        "Página de inicio.",
        "Nosotros.",
        "Servicios.",
        "Blog listo para publicar. Talitacumdacecol entrega los textos.",
        "Contacto, con formulario.",
        "Integración con WhatsApp.",
        "Integración con redes sociales.",
        "Google Maps.",
        "Diseño adaptable a celular, tableta y computador.",
        "Diseño con el logo, los colores y la información de Talitacumdacecol.",
        "Una ronda de ajustes después de la primera revisión.",
        "Panel administrativo.",
        "Configuración del sitio, del hosting y del certificado SSL.",
        "Capacitación de dos horas al publicar.",
    ]:
        add_bullet(doc, [(item, False)])
    add_bullet(doc, [("Optimización ", False), ("SEO", True), (" básica: títulos, descripciones y estructura de las páginas.", False)])

    add_subheading(doc, "2.2 Plan mensual de redes sociales")
    add_body(
        doc,
        [
            ("Valor: ", False),
            ("$1.000.000 COP al mes", True),
            (". Ese valor cubre las dos partes juntas: los 14 videos y las 16 imágenes. No se cobra cada parte por separado.", False),
        ],
        after=4,
    )
    add_body(
        doc,
        [
            ("Las piezas se publican en las redes que Talitacumdacecol tenga activas y a las que dé acceso, por ejemplo Instagram, Facebook, TikTok o las que indique. El material de base lo entrega ", False),
            ("EL CLIENTE", True),
            (" al empezar el mes: videos ya grabados y fotos. ", False),
            ("THE COMPANY DEV", True),
            (" no hace la grabación ni la sesión de fotos.", False),
        ],
    )

    add_body(doc, [("Videos: 14 en el mes", True)], before=4, after=2, align="left", keep_next=True)
    for parts in [
        [("Ayuda con el guion de cada video: estructura, mensaje y texto.", False)],
        [("Edición de cada video que Talitacumdacecol entregue grabado: corte, orden, texto en pantalla y música si ", False), ("EL CLIENTE", True), (" la entrega.", False)],
        [("Miniatura o portada del video, hecha con el material entregado.", False)],
        [("Subida de cada video a las redes activas.", False)],
        [("Texto corto de la publicación de cada video.", False)],
    ]:
        add_bullet(doc, parts)

    add_body(doc, [("Imágenes: 16 en el mes", True)], before=6, after=2, align="left", keep_next=True)
    for parts in [
        [("Edición de cada imagen: ajuste, formato para redes y texto sobre la pieza cuando haga falta.", False)],
        [("Texto del post de cada imagen.", False)],
        [("Subida de cada post a todas las redes activas de Talitacumdacecol.", False)],
    ]:
        add_bullet(doc, parts)

    add_body(
        doc,
        [
            ("Si en un mes Talitacumdacecol necesita más de 14 videos o más de 16 imágenes, ese adicional se pide y se cotiza. Lo que ya está publicado se mantiene.", False),
        ],
        before=4,
    )

    add_subheading(doc, "2.3 Campañas")
    add_body(
        doc,
        [
            ("La pauta, es decir el dinero que se paga a Instagram, Facebook u otra plataforma, lo asume ", False),
            ("EL CLIENTE", True),
            (". ", False),
            ("THE COMPANY DEV", True),
            (" cobra el trabajo de asesoría o de realización, no el presupuesto de la pauta.", False),
        ],
    )
    add_body(doc, [("Consultoría y asesoría: $50.000 COP por campaña", True)], before=4, after=2, align="left", keep_next=True)
    for parts in [
        [("Revisión del objetivo de la campaña.", False)],
        [("Definición del mensaje y del público.", False)],
        [("Recomendación de la red, del formato y del llamado a la acción.", False)],
        [("Orientación de cómo pautar, sin configurar todavía la campaña.", False)],
    ]:
        add_bullet(doc, parts)

    add_body(doc, [("Realización de cada campaña: $120.000 COP", True)], before=6, after=2, align="left", keep_next=True)
    add_body(
        doc,
        [("Este valor es por campaña y ya incluye la consultoría y la asesoría. Cubre:", False)],
        before=0,
        after=2,
        keep_next=True,
    )
    for parts in [
        [("Revisión del objetivo, del mensaje y del público.", False)],
        [("Ajuste de la pieza para pautar, con el material que entregue Talitacumdacecol.", False)],
        [("Configuración de la campaña en la plataforma que indique ", False), ("EL CLIENTE", True), (".", False)],
        [("Segmentación y publicación.", False)],
        [("Revisión básica de resultados al cierre de esa campaña.", False)],
    ]:
        add_bullet(doc, parts)

    add_heading(doc, "3. Información que debe entregar Talitacumdacecol")
    add_body(
        doc,
        [("Para el sitio web, al iniciar:", False)],
        keep_next=True,
    )
    for parts in [
        [("Hosting y dominio. Si todavía no los tiene, ", False), ("THE COMPANY DEV", True), (" asesora la compra. El costo lo asume ", False), ("EL CLIENTE", True), (".", False)],
        [("Manual de marca y logotipo en alta resolución.", False)],
        [("Fotografías.", False)],
        [("Historia de la empresa, servicios, teléfonos, correos y redes.", False)],
        [("Textos que quiera publicar.", False)],
    ]:
        add_bullet(doc, parts)
    add_body(
        doc,
        [("Para el plan de redes, al empezar cada mes:", False)],
        before=4,
        after=2,
        keep_next=True,
    )
    for parts in [
        [("Accesos de administrador a las redes activas.", False)],
        [("Los videos ya grabados, hasta 14 en el mes.", False)],
        [("Las fotos de base, hasta 16 en el mes.", False)],
        [("Datos de la marca: qué se quiere decir y a quién.", False)],
    ]:
        add_bullet(doc, parts)
    add_body(
        doc,
        [("Para una campaña:", False)],
        before=4,
        after=2,
        keep_next=True,
    )
    for parts in [
        [("Objetivo, fechas y presupuesto de pauta.", False)],
        [("Acceso a la cuenta publicitaria.", False)],
        [("Pieza o material de la campaña.", False)],
    ]:
        add_bullet(doc, parts)

    add_heading(doc, "4. Cronograma")
    add_body(
        doc,
        [
            ("El sitio web arranca cuando la propuesta está aceptada, está el primer pago y Talitacumdacecol entregó la información completa. Ese día ", False),
            ("THE COMPANY DEV", True),
            (" confirma la fecha de publicación.", False),
        ],
    )
    add_table(
        doc,
        ["ETAPA", "ACTIVIDAD", "MOMENTO"],
        [
            ["1", "Reunión inicial, recepción de la información y definición de la estructura del sitio.", "Al iniciar"],
            ["2", "Diseño y desarrollo del sitio con el contenido entregado.", "Después de recibir la información"],
            ["3", "Revisión, ronda de ajustes, capacitación y publicación.", "Al cierre"],
        ],
        [900, 6200, 1900],
    )
    add_body(
        doc,
        [
            ("El plan de redes corre por mes. Si los videos y las fotos llegan en los primeros días del mes, las 14 publicaciones de video y las 16 de imagen salen dentro de ese mes.", False),
        ],
        before=4,
    )

    add_heading(doc, "5. Términos y condiciones")
    terms = [
        [
            ("EL CLIENTE", True),
            (" entrega los accesos de hosting, dominio y redes que hagan falta para publicar.", False),
        ],
        [
            ("El sitio web inicia cuando la propuesta está aceptada y se ha hecho el primer pago.", False),
        ],
        [
            ("El plan de redes inicia con el pago del mes y con el material de ese mes.", False),
        ],
        [
            ("Si Talitacumdacecol entrega la información o el material después, las fechas se corren esos días.", False),
        ],
        [
            ("El hosting, el dominio, la pauta y los demás servicios externos los paga ", False),
            ("EL CLIENTE", True),
            (".", False),
        ],
        [
            ("Si después de aprobar el diseño pide un cambio de estructura, ", False),
            ("THE COMPANY DEV", True),
            (" lo revisa y lo cotiza.", False),
        ],
        [
            ("EL CLIENTE", True),
            (" puede empezar por el sitio, por las redes o por los dos. Más adelante puede pedir campañas, más piezas o funciones nuevas, y ese trabajo se cotiza cuando lo solicite.", False),
        ],
    ]
    for parts in terms:
        add_bullet(doc, parts)

    add_heading(doc, "6. Valor del servicio")
    add_subheading(doc, "Sitio web corporativo")
    add_body(doc, [("Valor: ", False), ("$800.000 COP", True)], before=2, after=4)
    add_body(
        doc,
        [
            ("Forma de pago: dos pagos iguales. El primero, de ", False),
            ("$400.000 COP", True),
            (", se hace al iniciar. El segundo, de ", False),
            ("$400.000 COP", True),
            (", se hace contra la entrega y la aprobación de ", False),
            ("EL CLIENTE", True),
            (".", False),
        ],
    )
    add_subheading(doc, "Plan de redes sociales")
    add_body(doc, [("Valor: ", False), ("$1.000.000 COP al mes", True)], before=2, after=4)
    add_body(
        doc,
        [
            ("Se paga al comenzar cada mes. Cubre los 14 videos, con ayuda de guion, edición y subida, y las 16 imágenes, con edición y subida del post a todas las redes activas.", False),
        ],
    )
    add_subheading(doc, "Campañas")
    add_body(
        doc,
        [
            ("Consultoría y asesoría: ", False),
            ("$50.000 COP", True),
            (" por campaña.", False),
        ],
        before=2,
        after=2,
    )
    add_body(
        doc,
        [
            ("Realización de la campaña: ", False),
            ("$120.000 COP", True),
            (" cada una. Este precio ya incluye la consultoría y la asesoría de esa campaña. La pauta no va dentro de este valor.", False),
        ],
    )
    add_body(doc, [("Medios de pago: Nequi, Daviplata o transferencia bancaria.", False)], before=4)
    add_body(
        doc,
        [
            ("El hosting y el dominio los paga ", False),
            ("EL CLIENTE", True),
            (". Si Talitacumdacecol no los tiene, ", False),
            ("THE COMPANY DEV", True),
            (" indica cómo adquirirlos. La renovación anual también queda a cargo de ", False),
            ("EL CLIENTE", True),
            (".", False),
        ],
    )

    add_heading(doc, "7. Entregables")
    add_body(doc, [("Del sitio web:", False)], before=2, after=2, align="left", keep_next=True)
    for item in [
        "Sitio publicado.",
        "Accesos al panel administrativo.",
        "Capacitación de dos horas al terminar.",
    ]:
        add_bullet(doc, [(item, False)])
    add_body(doc, [("Del mes de redes:", False)], before=4, after=2, align="left", keep_next=True)
    for item in [
        "14 videos editados y publicados, con el guion trabajado.",
        "16 imágenes editadas y publicadas en las redes activas.",
    ]:
        add_bullet(doc, [(item, False)])
    add_body(doc, [("De cada campaña realizada:", False)], before=4, after=2, align="left", keep_next=True)
    for item in [
        "Campaña configurada y publicada.",
        "Revisión básica de resultados al cierre.",
    ]:
        add_bullet(doc, [(item, False)])

    add_body(
        doc,
        [
            ("Quedo atento para revisar con Talitacumdacecol por dónde empezar y dejar la fecha de inicio.", False),
        ],
        before=8,
    )

    table = doc.add_table(rows=1, cols=2)
    table.autofit = True
    left, right = table.rows[0].cells
    for cell, lines in (
        (
            left,
            [
                ("THE COMPANY DEV", True),
                ("Jhon Javier Angulo Hilarion", False),
                ("C.C. 1024574523", False),
            ],
        ),
        (
            right,
            [
                ("EL CLIENTE", True),
                ("Talitacumdacecol", True),
            ],
        ),
    ):
        cell.text = ""
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcBorders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            element = OxmlElement(f"w:{edge}")
            element.set(qn("w:val"), "nil")
            tcBorders.append(element)
        tcPr.append(tcBorders)
        first = True
        for text, bold in lines:
            if first:
                paragraph = cell.paragraphs[0]
                first = False
            else:
                paragraph = cell.add_paragraph()
            set_paragraph_format(paragraph, before=0, after=0, align="left")
            run = paragraph.add_run(text)
            set_run_font(run, size=13, bold=bold, color=BLACK)
        space = cell.paragraphs[0]
        space.paragraph_format.space_before = Pt(18)

    add_contacts(doc)

    doc.core_properties.title = "Propuesta comercial Talitacumdacecol - The Company Dev"
    doc.core_properties.author = "The Company Dev"
    doc.core_properties.subject = "Sitio web corporativo, plan de redes y campañas"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
