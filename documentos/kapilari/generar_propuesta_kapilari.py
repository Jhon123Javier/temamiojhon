#!/usr/bin/env python3
"""Propuesta comercial KAPILARI con membrete de The Company Dev."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor, Twips

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUT = ROOT / "Propuesta The Company Dev - KAPILARI - Pagina Corporativa o Tienda Virtual.docx"

NAVY = RGBColor(0x00, 0x30, 0x60)
BLACK = RGBColor(0x1B, 0x33, 0x25)
GREEN = "02AF5D"
DARK = "1B3325"


def set_run_font(run, size=13, bold=False, color=None, name="Poppins"):
    run.bold = bold
    run.italic = False
    run.font.size = Pt(size)
    run.font.name = name
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), name)


def set_paragraph_format(paragraph, before=0, after=8, align="justify", keep_next=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = 1.08
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if align == "justify":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "center":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pPr = paragraph._p.get_or_add_pPr()
    if keep_next:
        if pPr.find(qn("w:keepNext")) is None:
            pPr.append(OxmlElement("w:keepNext"))
    if pPr.find(qn("w:widowControl")) is None:
        pPr.append(OxmlElement("w:widowControl"))


def add_runs(paragraph, parts, size=13, color=None):
    for text, bold in parts:
        run = paragraph.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color if color else BLACK)


def add_body(doc, parts, size=13, before=0, after=8, align="justify", keep_next=False):
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, before=before, after=after, align=align, keep_next=keep_next)
    add_runs(paragraph, parts, size=size)
    return paragraph


def add_heading(doc, text):
    return add_body(
        doc,
        [(text, True)],
        size=13,
        before=12,
        after=4,
        align="left",
        keep_next=True,
    )


def add_subheading(doc, text):
    return add_body(
        doc,
        [(text, True)],
        size=13,
        before=8,
        after=2,
        align="left",
        keep_next=True,
    )


def add_bullet(doc, parts, size=13):
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, before=1, after=2, align="left")
    fmt = paragraph.paragraph_format
    fmt.left_indent = Cm(0.75)
    fmt.first_line_indent = Cm(-0.4)
    add_runs(paragraph, [("•  ", False)] + parts, size=size)
    return paragraph


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_borders(cell, color="1B3325", sz="4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), sz)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        tcBorders.append(element)
    tcPr.append(tcBorders)


def set_cell_margins(cell, top=40, bottom=40, left=80, right=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def prevent_row_split(row, keep_with_next=False):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    if trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))
    if keep_with_next:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                pPr = paragraph._p.get_or_add_pPr()
                if pPr.find(qn("w:keepNext")) is None:
                    pPr.append(OxmlElement("w:keepNext"))


def write_cell(cell, text, bold=False, align="left", fill=None, size=12, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    set_paragraph_format(paragraph, before=0, after=0, align=align)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color or BLACK)
    set_cell_borders(cell)
    set_cell_margins(cell)
    if fill:
        shade_cell(cell, fill)
    # vertical align center
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    vAlign = OxmlElement("w:vAlign")
    vAlign.set(qn("w:val"), "center")
    tcPr.append(vAlign)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    total = sum(widths)
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    # fixed layout
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
    else:
        grid = OxmlElement("w:tblGrid")
        tblPr.addnext(grid)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        write_cell(cell, header, bold=True, align="center", fill=DARK, size=11, color=RGBColor(255, 255, 255))
        cell.width = Twips(widths[index])
    prevent_row_split(table.rows[0], keep_with_next=True)

    for r_index, row in enumerate(rows):
        for c_index, value in enumerate(row):
            align = "center" if c_index != 1 else "left"
            cell = table.rows[r_index + 1].cells[c_index]
            write_cell(cell, value, bold=False, align=align, fill="F7FBF8" if r_index % 2 else "FFFFFF", size=11)
            cell.width = Twips(widths[c_index])
        prevent_row_split(table.rows[r_index + 1], keep_with_next=r_index < len(rows) - 1)

    # space after table
    spacer = doc.add_paragraph()
    set_paragraph_format(spacer, before=0, after=2, align="left")
    return table


def add_background(header, image_path):
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_format(paragraph, before=0, after=0, align="left")
    run = paragraph.add_run()
    set_run_font(run, size=1)
    run.add_picture(image_path, width=Inches(8.5), height=Inches(11))
    drawing = run._element.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))
    extent = inline.find(qn("wp:extent"))
    cx = extent.get("cx")
    cy = extent.get("cy")
    doc_pr = inline.find(qn("wp:docPr"))
    graphic = inline.find(qn("a:graphic"))
    anchor = parse_xml(
        f'''
        <wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
                   xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
                   distT="0" distB="0" distL="0" distR="0"
                   simplePos="0" relativeHeight="0" behindDoc="1"
                   locked="0" layoutInCell="1" allowOverlap="1">
          <wp:simplePos x="0" y="0"/>
          <wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>
          <wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>
          <wp:extent cx="{cx}" cy="{cy}"/>
          <wp:effectExtent l="0" t="0" r="0" b="0"/>
          <wp:wrapNone/>
          <wp:docPr id="{doc_pr.get("id")}" name="membrete"/>
          <wp:cNvGraphicFramePr>
            <a:graphicFrameLocks xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" noChangeAspect="1"/>
          </wp:cNvGraphicFramePr>
        </wp:anchor>
        '''
    )
    anchor.append(graphic)
    drawing.replace(inline, anchor)


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.72)
    section.bottom_margin = Inches(1.08)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.15)
    section.footer_distance = Inches(0.15)
    section.different_first_page_header_footer = False
    sectPr = section._sectPr
    # one header for every page
    for tag in ("w:titlePg",):
        node = sectPr.find(qn(tag))
        if node is not None:
            sectPr.remove(node)


def add_label_line(doc, label, value):
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, before=0, after=1, align="left")
    label_run = paragraph.add_run(label)
    set_run_font(label_run, size=13, bold=True, color=NAVY)
    value_run = paragraph.add_run(value)
    set_run_font(value_run, size=13, bold=False, color=BLACK)
    # tab-like spacing via two runs; pad label
    return paragraph


def add_contacts(doc):
    contacts = [
        (ASSETS / "icon_telefono.png", "3107862438 - 3118620020"),
        (ASSETS / "icon_correo.png", "info@thecompanydev.com"),
        (ASSETS / "icon_web.png", "www.thecompanydev.com"),
    ]
    for index, (icon, text) in enumerate(contacts):
        paragraph = doc.add_paragraph()
        set_paragraph_format(paragraph, before=200 if index == 0 else 1, after=0, align="right")
        if index < len(contacts) - 1:
            pPr = paragraph._p.get_or_add_pPr()
            if pPr.find(qn("w:keepNext")) is None:
                pPr.append(OxmlElement("w:keepNext"))
        icon_run = paragraph.add_run()
        icon_run.add_picture(str(icon), width=Inches(0.2), height=Inches(0.2))
        text_run = paragraph.add_run("  " + text)
        set_run_font(text_run, size=11, bold=False, color=BLACK)


def build():
    doc = Document()
    configure_section(doc.sections[0])
    add_background(doc.sections[0].header, str(ASSETS / "membrete_fondo_carta_sin_contacto.png"))

    normal = doc.styles["Normal"]
    normal.font.name = "Poppins"
    normal.font.size = Pt(13)
    normal.font.color.rgb = BLACK
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), "Poppins")

    add_label_line(doc, "Fecha     ", "9 de octubre de 2026")
    add_label_line(doc, "Asunto   ", "Propuesta de Prestación de Servicios")
    add_label_line(doc, "Para      ", "KAPILARI")

    add_body(
        doc,
        [("PROPUESTA COMERCIAL", True)],
        size=14,
        before=10,
        after=0,
        align="left",
        keep_next=True,
    )
    add_body(
        doc,
        [("OPCIÓN A: PÁGINA CORPORATIVA    ·    OPCIÓN B: TIENDA VIRTUAL", True)],
        size=13,
        before=0,
        after=8,
        align="left",
    ).runs[0].font.color.rgb = NAVY
    # color was set on font after add; reapply via rPr already done if RGB assigned
    # The helper already used BLACK. Fix the title color properly.
    title = doc.paragraphs[-1]
    for run in title.runs:
        set_run_font(run, size=13, bold=True, color=NAVY)

    add_body(
        doc,
        [
            ("Por medio de la presente, ", False),
            ("THE COMPANY DEV", True),
            (" presenta a KAPILARI la propuesta para el diseño, desarrollo e implementación de su presencia digital. ", False),
            ("EL CLIENTE", True),
            (" elige una sola opción. No son etapas que se hagan una después de la otra, y esta propuesta no incluye plan de crecimiento ni administración mensual.", False),
        ],
    )
    add_body(
        doc,
        [
            ("Opción A. Página corporativa. ", True),
            ("Sirve para presentar la empresa, sus servicios y sus datos de contacto. Valor: ", False),
            ("$700.000 COP", True),
            (". Tiempo máximo de entrega: ", False),
            ("1 semana", True),
            (".", False),
        ],
    )
    add_body(
        doc,
        [
            ("Opción B. Tienda virtual. ", True),
            ("Sirve para mostrar productos, recibir pedidos y vender en línea. Valor: ", False),
            ("$1.400.000 COP", True),
            (". Tiempo máximo de entrega: ", False),
            ("2 semanas", True),
            (".", False),
        ],
    )
    add_body(
        doc,
        [
            ("Esos plazos se cuentan en días hábiles y aplican únicamente si KAPILARI entrega la información completa en el momento en que ", False),
            ("THE COMPANY DEV", True),
            (" la solicite. Si la información llega después, el plazo se corre los días hábiles del retraso.", False),
        ],
    )

    add_heading(doc, "1. Objetivo del servicio")
    add_body(
        doc,
        [
            ("Diseñar, desarrollar y publicar la opción que elija ", False),
            ("EL CLIENTE", True),
            (": una página corporativa o una tienda virtual. El sitio queda administrable, se ve bien en celular y en computador, y usa la identidad de KAPILARI.", False),
        ],
    )

    add_heading(doc, "2. Alcance del servicio")
    add_subheading(doc, "Opción A. Página corporativa")
    add_body(
        doc,
        [
            ("Valor del proyecto: ", False),
            ("$700.000 COP", True),
            (". Tiempo máximo: ", False),
            ("1 semana", True),
            (", si la información llega a tiempo.", False),
        ],
        after=4,
    )
    add_body(
        doc,
        [
            ("La página presenta quién es KAPILARI, qué ofrece y cómo contactarla. Es el canal institucional de la empresa.", False),
        ],
    )

    add_subheading(doc, "Opción B. Tienda virtual")
    add_body(
        doc,
        [
            ("Valor del proyecto: ", False),
            ("$1.400.000 COP", True),
            (". Tiempo máximo: ", False),
            ("2 semanas", True),
            (", si la información llega a tiempo.", False),
        ],
        after=4,
    )
    add_body(
        doc,
        [
            ("La tienda permite publicar el catálogo, armar el pedido y gestionarlo desde un panel. La pasarela de pago no está incluida.", False),
        ],
    )

    add_subheading(doc, "2.1 Funcionalidades")
    add_body(doc, [("Página corporativa", True)], before=2, after=2, align="left", keep_next=True)
    for item in [
        "Página de inicio.",
        "Nosotros.",
        "Servicios.",
        "Blog, listo para publicar. La redacción de artículos no está incluida.",
        "Contacto, con formulario.",
        "Integración con WhatsApp.",
        "Integración con redes sociales.",
        "Google Maps.",
        "Diseño adaptable a celular, tableta y computador.",
        "Panel administrativo.",
    ]:
        add_bullet(doc, [(item, False)])
    add_bullet(doc, [("Optimización ", False), ("SEO", True), (" básica.", False)])

    add_body(doc, [("Tienda virtual", True)], before=6, after=2, align="left", keep_next=True)
    for item in [
        "Catálogo de productos.",
        "Categorías y variantes.",
        "Carrito de compras y checkout.",
        "Gestión de pedidos.",
        "Métodos de envío.",
        "Gestión de inventario básica.",
        "Panel administrativo.",
        "Carga inicial de los productos que KAPILARI entregue al comenzar, con foto, nombre, descripción, categoría y precio.",
        "Diseño adaptable a celular, tableta y computador.",
        "Integración con WhatsApp.",
    ]:
        add_bullet(doc, [(item, False)])
    add_bullet(
        doc,
        [("Integración con pasarela de pago. La pasarela no está incluida en este valor.", False)],
    )

    add_subheading(doc, "2.2 Diseño")
    add_body(
        doc,
        [
            ("El sitio de la opción elegida se diseña con la identidad de KAPILARI: logo, colores y forma de mostrar la información. Dentro del plazo hay una ronda de ajustes.", False),
        ],
    )
    for item in [
        "Diseño a partir del manual de marca y el logotipo que entregue KAPILARI.",
        "Navegación clara.",
        "Buena lectura en computador y en celular.",
        "Una ronda de ajustes después de la primera revisión.",
    ]:
        add_bullet(doc, [(item, False)])

    add_subheading(doc, "2.3 Estructura")
    add_body(doc, [("Página corporativa: ", True), ("Inicio, Nosotros, Servicios, Blog y Contacto.", False)], before=2, after=2)
    add_body(
        doc,
        [
            ("Tienda virtual: ", True),
            ("Inicio, Tienda, Categorías, Producto, Carrito, Checkout, Nosotros, Contacto y políticas legales. Los textos legales los entrega KAPILARI.", False),
        ],
    )

    add_subheading(doc, "2.4 Configuración e información que debe entregar KAPILARI")
    add_body(
        doc,
        [
            ("THE COMPANY DEV", True),
            (" configura el sitio, el hosting, el certificado SSL y los formularios, hace la optimización inicial y da una capacitación básica al entregar. En la tienda virtual también configura el catálogo, los métodos de envío y las pruebas del pedido.", False),
        ],
    )
    add_body(
        doc,
        [
            ("Para iniciar, y para que el plazo de 1 semana o de 2 semanas se cumpla, KAPILARI debe entregar en ese momento:", False),
        ],
        keep_next=True,
    )
    for item in [
        "Hosting y dominio. Si todavía no los tiene, THE COMPANY DEV asesora la compra. El costo lo asume EL CLIENTE.",
        "Manual de marca y logotipo en alta resolución.",
        "Fotografías.",
        "Historia de la empresa e información institucional.",
        "Teléfonos, correos y redes sociales.",
        "Textos de los servicios, si elige la página corporativa.",
    ]:
        if item.startswith("Hosting"):
            add_bullet(
                doc,
                [
                    ("Hosting y dominio. Si todavía no los tiene, ", False),
                    ("THE COMPANY DEV", True),
                    (" asesora la compra. El costo lo asume ", False),
                    ("EL CLIENTE", True),
                    (".", False),
                ],
            )
        else:
            add_bullet(doc, [(item, False)])
    add_body(
        doc,
        [("Si elige la tienda virtual, también debe entregar al inicio:", False)],
        before=4,
        after=2,
        keep_next=True,
    )
    for item in [
        "Fotografías de los productos.",
        "Nombre, descripción y categoría de cada producto.",
        "Precios.",
        "Inventario inicial.",
        "Textos de las políticas legales.",
    ]:
        add_bullet(doc, [(item, False)])

    add_subheading(doc, "2.5 Soporte durante la entrega")
    add_body(
        doc,
        [
            ("Durante el desarrollo, la atención es por WhatsApp y por correo electrónico, de lunes a viernes, de 9:00 a.m. a 5:00 p.m. El tiempo máximo de respuesta es de 24 horas hábiles.", False),
        ],
    )
    add_body(
        doc,
        [
            ("La administración mensual, el mantenimiento continuo, las copias de seguridad periódicas y el plan de crecimiento no están incluidos.", False),
        ],
    )

    add_heading(doc, "3. Cronograma")
    add_body(
        doc,
        [
            ("Los tiempos de abajo son el máximo de cada opción. Cuentan desde el primer pago y desde el momento en que KAPILARI entrega la información completa. Si esa información no llega cuando se pide, el cronograma se corre.", False),
        ],
    )
    add_body(doc, [("Opción A. Página corporativa: máximo 1 semana", True)], before=2, after=4, align="left", keep_next=True)
    add_table(
        doc,
        ["ETAPA", "ACTIVIDAD", "TIEMPO"],
        [
            ["1", "Recepción de la información, reunión inicial y definición de la estructura.", "Día 1"],
            ["2", "Diseño y desarrollo de la página corporativa con el contenido entregado.", "Días 2 a 4"],
            ["3", "Revisión, ronda de ajustes, capacitación, publicación y entrega.", "Día 5"],
        ],
        [900, 6200, 1600],
    )
    add_body(doc, [("Opción B. Tienda virtual: 2 semanas", True)], before=4, after=4, align="left", keep_next=True)
    add_table(
        doc,
        ["ETAPA", "ACTIVIDAD", "TIEMPO"],
        [
            ["1", "Recepción de la información, reunión inicial, estructura y diseño de la tienda.", "Semana 1"],
            ["2", "Desarrollo de la tienda, carga inicial de productos y configuración.", "Semana 1"],
            ["3", "Pruebas, ronda de ajustes, capacitación, publicación y entrega.", "Semana 2"],
        ],
        [900, 6200, 1600],
    )

    add_heading(doc, "4. Términos y condiciones")
    terms = [
        [
            ("EL CLIENTE", True),
            (" entrega los accesos de hosting, dominio y demás cuentas que hagan falta para publicar el sitio.", False),
        ],
        [
            ("El trabajo inicia cuando la propuesta está aceptada y se ha hecho el primer pago.", False),
        ],
        [
            ("Lo que no esté escrito en esta propuesta se cotiza aparte.", False),
        ],
        [
            ("El plazo máximo es de 1 semana para la página corporativa y de 2 semanas para la tienda virtual. Esos plazos aplican si KAPILARI entrega la información completa en el momento en que se le solicita. Si la entrega se retrasa, el plazo se corre en los días hábiles de ese retraso.", False),
        ],
        [
            ("El hosting, el dominio, la pasarela de pago y los demás servicios externos los paga ", False),
            ("EL CLIENTE", True),
            (".", False),
        ],
        [
            ("Los cambios de estructura pedidos después de aprobar el diseño pueden tener un costo adicional.", False),
        ],
        [
            ("Esta propuesta no incluye plan de crecimiento, administración mensual ni mantenimiento continuo.", False),
        ],
        [
            ("EL CLIENTE", True),
            (" elige la Opción A o la Opción B. Los dos valores no se suman, porque no se contratan las dos opciones dentro de esta propuesta.", False),
        ],
    ]
    for parts in terms:
        add_bullet(doc, parts)

    add_heading(doc, "5. Valor del servicio")
    add_subheading(doc, "Opción A. Página corporativa")
    add_body(doc, [("Valor: ", False), ("$700.000 COP", True)], before=2, after=4)
    add_body(
        doc,
        [
            ("Forma de pago: dos pagos iguales. El primero, de ", False),
            ("$350.000 COP", True),
            (", se hace al iniciar, para arrancar el desarrollo. El segundo, de ", False),
            ("$350.000 COP", True),
            (", se hace contra la entrega final y la aprobación de ", False),
            ("EL CLIENTE", True),
            (".", False),
        ],
    )
    add_subheading(doc, "Opción B. Tienda virtual")
    add_body(doc, [("Valor: ", False), ("$1.400.000 COP", True)], before=2, after=4)
    add_body(
        doc,
        [
            ("Forma de pago: dos pagos iguales. El primero, de ", False),
            ("$700.000 COP", True),
            (", se hace al iniciar, para arrancar el desarrollo. El segundo, de ", False),
            ("$700.000 COP", True),
            (", se hace contra la entrega final y la aprobación de ", False),
            ("EL CLIENTE", True),
            (".", False),
        ],
    )
    add_body(doc, [("Medios de pago: Nequi, Daviplata o transferencia bancaria.", False)], before=4)
    add_body(
        doc,
        [
            ("El hosting y el dominio no están incluidos en estos valores. Si KAPILARI no los tiene, ", False),
            ("THE COMPANY DEV", True),
            (" indica cómo adquirirlos. La renovación anual queda a cargo de ", False),
            ("EL CLIENTE", True),
            (".", False),
        ],
    )

    break_paragraph = doc.add_paragraph()
    set_paragraph_format(break_paragraph, before=0, after=0, align="left")
    break_run = break_paragraph.add_run()
    set_run_font(break_run, size=1)
    break_run.add_break(WD_BREAK.PAGE)

    add_heading(doc, "6. Entregables")
    for item in [
        "Sitio publicado de la opción elegida.",
        "Accesos al panel administrativo.",
        "Capacitación de dos horas al terminar la entrega.",
    ]:
        add_bullet(doc, [(item, False)])

    add_body(
        doc,
        [
            ("Quedo atento para revisar con KAPILARI cuál de las dos opciones se ajusta a lo que necesita y dejar la fecha de inicio.", False),
        ],
        before=8,
    )

    # Firmas
    table = doc.add_table(rows=1, cols=2)
    table.autofit = True
    left, right = table.rows[0].cells
    for cell, lines in (
        (
            left,
            [
                ("THE COMPANY DEV", True),
                ("Jhon Javier Angulo Hilarion", False),
                ("C.C. 1024574523", False),
            ],
        ),
        (
            right,
            [
                ("EL CLIENTE", True),
                ("KAPILARI", True),
            ],
        ),
    ):
        cell.text = ""
        # remove borders
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcBorders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            element = OxmlElement(f"w:{edge}")
            element.set(qn("w:val"), "nil")
            tcBorders.append(element)
        tcPr.append(tcBorders)
        first = True
        for text, bold in lines:
            if first:
                paragraph = cell.paragraphs[0]
                first = False
            else:
                paragraph = cell.add_paragraph()
            set_paragraph_format(paragraph, before=0, after=0, align="left")
            run = paragraph.add_run(text)
            set_run_font(run, size=13, bold=bold, color=BLACK)
        # signature space
        space = cell.paragraphs[0]
        space.paragraph_format.space_before = Pt(18)

    add_contacts(doc)

    core = doc.part.element
    # document title
    doc.core_properties.title = "Propuesta comercial KAPILARI - The Company Dev"
    doc.core_properties.author = "The Company Dev"
    doc.core_properties.subject = "Opción A página corporativa u opción B tienda virtual"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
