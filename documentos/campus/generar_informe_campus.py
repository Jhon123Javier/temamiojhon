#!/usr/bin/env python3
"""Informe de perfil Campus by TCD con membrete de The Company Dev."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor, Twips

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUT = ROOT / "Informe The Company Dev - Perfil Campus by TCD.docx"

NAVY = RGBColor(0x00, 0x30, 0x60)
BLACK = RGBColor(0x1B, 0x33, 0x25)
GREEN = "02AF5D"
DARK = "1B3325"


def set_run_font(run, size=13, bold=False, color=None, name="Poppins"):
    run.bold = bold
    run.italic = False
    run.font.size = Pt(size)
    run.font.name = name
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), name)


def set_paragraph_format(paragraph, before=0, after=8, align="justify", keep_next=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = 1.08
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if align == "justify":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "center":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pPr = paragraph._p.get_or_add_pPr()
    if keep_next:
        if pPr.find(qn("w:keepNext")) is None:
            pPr.append(OxmlElement("w:keepNext"))
    if pPr.find(qn("w:widowControl")) is None:
        pPr.append(OxmlElement("w:widowControl"))


KEYWORDS = (
    "EL CLIENTE",
    "EL CONTRATISTA",
    "THE COMPANY DEV",
    "RECUPERACO",
    "RECOROUTE",
    "RecuperaCo",
    "RecoRoute",
    "SEO",
    "Seo",
)


def split_keywords(text):
    import re

    pattern = re.compile("(" + "|".join(re.escape(word) for word in KEYWORDS) + ")")
    return [piece for piece in pattern.split(text) if piece]


def add_runs(paragraph, parts, size=13, color=None):
    for text, bold in parts:
        for piece in split_keywords(text):
            run = paragraph.add_run(piece)
            set_run_font(
                run,
                size=size,
                bold=bold or piece in KEYWORDS,
                color=color if color else BLACK,
            )


def add_body(doc, parts, size=13, before=0, after=8, align="justify", keep_next=False):
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, before=before, after=after, align=align, keep_next=keep_next)
    add_runs(paragraph, parts, size=size)
    return paragraph


def add_heading(doc, text):
    return add_body(
        doc,
        [(text, True)],
        size=13,
        before=12,
        after=4,
        align="left",
        keep_next=True,
    )


def add_subheading(doc, text):
    return add_body(
        doc,
        [(text, True)],
        size=13,
        before=8,
        after=2,
        align="left",
        keep_next=True,
    )


def add_bullet(doc, parts, size=13):
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, before=1, after=2, align="left")
    fmt = paragraph.paragraph_format
    fmt.left_indent = Cm(0.75)
    fmt.first_line_indent = Cm(-0.4)
    add_runs(paragraph, [("•  ", False)] + parts, size=size)
    return paragraph


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_borders(cell, color="1B3325", sz="4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), sz)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        tcBorders.append(element)
    tcPr.append(tcBorders)


def set_cell_margins(cell, top=40, bottom=40, left=80, right=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def prevent_row_split(row, keep_with_next=False):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    if trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))
    if keep_with_next:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                pPr = paragraph._p.get_or_add_pPr()
                if pPr.find(qn("w:keepNext")) is None:
                    pPr.append(OxmlElement("w:keepNext"))


def write_cell(cell, text, bold=False, align="left", fill=None, size=12, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    set_paragraph_format(paragraph, before=0, after=0, align=align)
    for piece in split_keywords(text):
        run = paragraph.add_run(piece)
        set_run_font(run, size=size, bold=bold or piece in KEYWORDS, color=color or BLACK)
    set_cell_borders(cell)
    set_cell_margins(cell)
    if fill:
        shade_cell(cell, fill)
    # vertical align center
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    vAlign = OxmlElement("w:vAlign")
    vAlign.set(qn("w:val"), "center")
    tcPr.append(vAlign)


def add_table(doc, headers, rows, widths, aligns=None, size=10, header_size=10):
    if aligns is None:
        aligns = ["left"] * len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    total = sum(widths)
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
    else:
        grid = OxmlElement("w:tblGrid")
        tblPr.addnext(grid)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        write_cell(
            cell,
            header,
            bold=True,
            align="center",
            fill=DARK,
            size=header_size,
            color=RGBColor(255, 255, 255),
        )
        cell.width = Twips(widths[index])
    prevent_row_split(table.rows[0], keep_with_next=True)

    for r_index, row in enumerate(rows):
        fill = "F7FBF8" if r_index % 2 else "FFFFFF"
        for c_index, value in enumerate(row):
            cell = table.rows[r_index + 1].cells[c_index]
            write_cell(cell, value, bold=False, align=aligns[c_index], fill=fill, size=size)
            cell.width = Twips(widths[c_index])
        prevent_row_split(table.rows[r_index + 1])

    spacer = doc.add_paragraph()
    set_paragraph_format(spacer, before=0, after=2, align="left")
    return table


def clear_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    tblPr.append(borders)


def set_table_indent(table, indent):
    tblPr = table._tbl.tblPr
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), str(indent))
    tblInd.set(qn("w:type"), "dxa")


def add_note(doc, text):
    return add_body(doc, [(text, False)], size=11, before=0, after=6, align="left")


def add_contacts(doc):
    contacts = [
        (ASSETS / "icon_telefono.png", "3107862438 - 3118620020"),
        (ASSETS / "icon_correo.png", "info@thecompanydev.com"),
        (ASSETS / "icon_web.png", "www.thecompanydev.com"),
    ]
    anchor = doc.add_paragraph()
    set_paragraph_format(anchor, before=250, after=0, align="right", keep_next=True)

    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    table.autofit = False
    widths = [420, 3480]
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    set_table_indent(table, 9648 - sum(widths))
    clear_table_borders(table)
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
    else:
        grid = OxmlElement("w:tblGrid")
        tblPr.addnext(grid)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for index, (icon, text) in enumerate(contacts):
        row = table.rows[index]
        prevent_row_split(row, keep_with_next=index < len(contacts) - 1)
        icon_cell = row.cells[0]
        text_cell = row.cells[1]
        for cell in (icon_cell, text_cell):
            set_cell_margins(cell, top=20, bottom=20, left=40, right=40)
        icon_cell.text = ""
        paragraph = icon_cell.paragraphs[0]
        set_paragraph_format(paragraph, before=0, after=0, align="right")
        run = paragraph.add_run()
        run.add_picture(str(icon), width=Inches(0.2), height=Inches(0.2))
        text_cell.text = ""
        paragraph = text_cell.paragraphs[0]
        set_paragraph_format(paragraph, before=0, after=0, align="left")
        text_run = paragraph.add_run(text)
        set_run_font(text_run, size=11, bold=False, color=BLACK)


def add_background(header, image_path):
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_format(paragraph, before=0, after=0, align="left")
    run = paragraph.add_run()
    set_run_font(run, size=1)
    run.add_picture(image_path, width=Inches(8.5), height=Inches(11))
    drawing = run._element.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))
    extent = inline.find(qn("wp:extent"))
    cx = extent.get("cx")
    cy = extent.get("cy")
    doc_pr = inline.find(qn("wp:docPr"))
    graphic = inline.find(qn("a:graphic"))
    anchor = parse_xml(
        f'''
        <wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
                   xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
                   distT="0" distB="0" distL="0" distR="0"
                   simplePos="0" relativeHeight="0" behindDoc="1"
                   locked="0" layoutInCell="1" allowOverlap="1">
          <wp:simplePos x="0" y="0"/>
          <wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>
          <wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>
          <wp:extent cx="{cx}" cy="{cy}"/>
          <wp:effectExtent l="0" t="0" r="0" b="0"/>
          <wp:wrapNone/>
          <wp:docPr id="{doc_pr.get("id")}" name="membrete"/>
          <wp:cNvGraphicFramePr>
            <a:graphicFrameLocks xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" noChangeAspect="1"/>
          </wp:cNvGraphicFramePr>
        </wp:anchor>
        '''
    )
    anchor.append(graphic)
    drawing.replace(inline, anchor)


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.72)
    section.bottom_margin = Inches(1.08)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.15)
    section.footer_distance = Inches(0.15)
    section.different_first_page_header_footer = False
    sectPr = section._sectPr
    # one header for every page
    for tag in ("w:titlePg",):
        node = sectPr.find(qn(tag))
        if node is not None:
            sectPr.remove(node)


def add_label_line(doc, label, value):
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, before=0, after=1, align="left")
    label_run = paragraph.add_run(label)
    set_run_font(label_run, size=13, bold=True, color=NAVY)
    for piece in split_keywords(value):
        value_run = paragraph.add_run(piece)
        set_run_font(value_run, size=13, bold=piece in KEYWORDS, color=BLACK)
    # tab-like spacing via two runs; pad label
    return paragraph




def build():
    doc = Document()
    configure_section(doc.sections[0])
    add_background(doc.sections[0].header, str(ASSETS / "membrete_fondo_carta_sin_contacto.png"))

    normal = doc.styles["Normal"]
    normal.font.name = "Poppins"
    normal.font.size = Pt(13)
    normal.font.color.rgb = BLACK
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), "Poppins")

    add_label_line(doc, "Fecha     ", "13 de octubre de 2026")
    add_label_line(
        doc,
        "Asunto   ",
        "Perfil del servicio Campus by TCD: identidad, alcance, mercado y precios",
    )
    add_label_line(doc, "Para      ", "Directivos de THE COMPANY DEV")

    add_body(
        doc,
        [("INFORME DE PERFIL DE SERVICIO", True)],
        size=13,
        before=12,
        after=0,
        align="left",
        keep_next=True,
    )
    add_body(
        doc,
        [("CAMPUS BY TCD", True)],
        size=14,
        before=2,
        after=0,
        align="left",
        keep_next=True,
    ).runs[0].font.color.rgb = NAVY
    add_body(
        doc,
        [("Gestión escolar e intervención temprana", True)],
        size=13,
        before=0,
        after=2,
        align="left",
        keep_next=True,
    )
    add_body(
        doc,
        [("“La matrícula se cuida cuando el riesgo se ve a tiempo.”", False)],
        size=13,
        before=0,
        after=8,
        align="left",
    )

    add_heading(doc, "RESUMEN EJECUTIVO")
    add_body(
        doc,
        [
            (
                "Este informe presenta a la directiva el perfil de Campus by TCD, el ecosistema de ",
                False,
            ),
            ("THE COMPANY DEV", True),
            (
                " para la gestión escolar y la intervención temprana. Reúne la identidad del servicio, el problema que atiende, el alcance de los cinco módulos, los precios que hoy publican otros proveedores en Colombia y una propuesta de modalidades y tarifas para aprobación.",
                False,
            ),
        ],
    )
    add_body(
        doc,
        [
            (
                "La propuesta comercial recomienda tres modalidades, con la configuración y la capacitación incluidas: suscripción mensual, anualidad y licencia vitalicia. Los valores de Campus que aparecen desde la sección 11 son una recomendación de este informe. Los de la sección 9 son cifras leídas en páginas y documentos públicos el 13 de octubre de 2026.",
                False,
            )
        ],
    )
    add_table(
        doc,
        ["5 módulos", "4 portales", "3 modalidades", "10 a 14 semanas"],
        [
            [
                "Admisiones, académico, alertas, refuerzo y rectoría.",
                "Rector, docente, acudiente y estudiante.",
                "Mensual, anualidad y licencia vitalicia.",
                "Puesta en marcha, contadas desde la entrega de la información.",
            ]
        ],
        [2400, 2400, 2400, 2400],
        aligns=["left", "left", "left", "left"],
        size=10,
    )

    add_heading(doc, "1. QUÉ ES CAMPUS BY TCD")
    add_body(
        doc,
        [
            (
                "Campus by TCD es un software de gestión escolar para colegios privados. Une la matrícula digital, el libro de calificaciones y el control de asistencia con un algoritmo que detecta bajo rendimiento e inasistencia. Cuando una regla se cumple, el sistema avisa al acudiente por WhatsApp, correo y el portal, y abre de inmediato la cápsula de refuerzo de la competencia que el estudiante no superó.",
                False,
            )
        ],
    )
    add_body(
        doc,
        [
            (
                "El cierre de cada alerta queda en una bitácora con fecha, hora, acudiente notificado y acción tomada. Esa evidencia respalda los planes de apoyo del Decreto 1290 de 2009 y le da al colegio la trazabilidad para cuidar la permanencia de sus estudiantes y sostener sus estándares de calidad.",
                False,
            )
        ],
    )

    add_heading(doc, "2. MISIÓN")
    add_body(
        doc,
        [
            (
                "En Campus by TCD convertimos la administración académica en un sistema de retención. Damos al colegio privado las herramientas para matricular, evaluar y pasar asistencia, y para actuar el mismo día en que un estudiante entra en riesgo. Así el acudiente se entera a tiempo, el docente tiene una ruta de nivelación lista y rectoría conserva la evidencia del acompañamiento.",
                False,
            )
        ],
    )

    add_heading(doc, "3. VISIÓN")
    add_body(
        doc,
        [
            (
                "En 2030, Campus by TCD será la plataforma de referencia en Colombia para los colegios privados que quieren ver el riesgo académico antes de perder la matrícula. Queremos ser reconocidos porque su bitácora sirve de soporte en auditorías, certificaciones de calidad y seguimiento del Decreto 1290, y porque cada alerta termina en un refuerzo que el estudiante puede completar.",
                False,
            )
        ],
    )

    add_heading(doc, "4. VALORES")
    add_table(
        doc,
        ["Valor", "Qué significa en el servicio"],
        [
            ["Oportunidad", "El riesgo se señala mientras la nivelación todavía alcanza."],
            ["Trazabilidad", "Cada alerta guarda fecha, hora, acudiente y acción tomada."],
            ["Acompañamiento", "El acudiente confirma que se enteró y ve la ruta de solución."],
            ["Claridad", "El docente registra notas y asistencia en una planilla de pocos clics."],
            ["Cuidado de la matrícula", "La evidencia del apoyo pedagógico protege la permanencia."],
        ],
        [2800, 6800],
        aligns=["left", "left"],
    )

    add_heading(doc, "5. PROBLEMA QUE RESUELVE")
    add_body(
        doc,
        [
            (
                "Cuando las notas viven en un archivo, la inasistencia se cuenta al cierre del periodo y el acudiente se entera en la entrega de boletines, el colegio llega tarde al apoyo. Campus by TCD ordena ese recorrido así:",
                False,
            )
        ],
    )
    add_table(
        doc,
        ["Hoy, sin Campus by TCD", "Con Campus by TCD"],
        [
            [
                "La preinscripción y los documentos llegan por correo o en papel.",
                "Formulario público, carga digital y ficha acumulativa del estudiante.",
            ],
            [
                "Las notas se arman al cierre, con una escala distinta en cada archivo.",
                "Libro de calificaciones por porcentajes, periodos, competencias o logros del SIEE.",
            ],
            [
                "La inasistencia se descubre cuando ya se acumuló.",
                "Registro diario o por clase, con alerta al porcentaje de fallas que defina el colegio.",
            ],
            [
                "El acudiente se entera en la reunión o en el boletín.",
                "Aviso inmediato por WhatsApp, correo y notificación en el portal de padres.",
            ],
            [
                "El plan de apoyo del Decreto 1290 se reconstruye a mano.",
                "Bitácora automática de fecha, hora, acudiente notificado y acción tomada.",
            ],
            [
                "El refuerzo depende de que el docente alcance a prepararlo.",
                "Cápsula de la competencia no superada, con guía, video y quiz de nivelación.",
            ],
            [
                "Rectoría pide el listado de riesgo para la reunión.",
                "Semáforo por grados y materias, exportable para auditoría y calidad.",
            ],
        ],
        [4800, 4800],
        aligns=["left", "left"],
    )

    add_heading(doc, "6. MERCADO OBJETIVO")
    add_bullet(
        doc,
        [
            (
                "Colegios privados de educación básica y media que necesitan matrícula digital, calificaciones, asistencia y comunicación con las familias.",
                False,
            )
        ],
    )
    add_bullet(
        doc,
        [
            (
                "Instituciones que deben dejar evidencia de los planes de apoyo y de la evaluación, en el marco del Decreto 1290 de 2009 y de su propio SIEE.",
                False,
            )
        ],
    )
    add_bullet(
        doc,
        [
            (
                "Rectorías que miden la deserción posible y quieren una alerta antes de perder la matrícula del siguiente año.",
                False,
            )
        ],
    )
    add_body(
        doc,
        [
            (
                "El comprador de esta versión es el colegio privado. Las cifras de colegios oficiales que aparecen en la sección 9 sirven solo como referencia de una licencia académica anual. Campus by TCD puede crecer después, por cotización, hacia cartera, SIMAT o varias sedes cuando un cliente lo pida.",
                False,
            )
        ],
    )

    add_heading(doc, "7. CÓMO FUNCIONA")
    add_body(
        doc,
        [
            (
                "Campus by TCD conecta cuatro portales. El acceso de cada persona se controla con usuarios, roles y permisos: administrador, rector, coordinador, docente, estudiante y padre o acudiente.",
                False,
            )
        ],
    )
    add_table(
        doc,
        ["Portal", "Qué hace en Campus by TCD"],
        [
            [
                "Rector o director académico",
                "Ve el semáforo de salud, la tasa de deserción potencial y el avance de los refuerzos. Exporta la evidencia para auditorías y certificaciones de calidad.",
            ],
            [
                "Docente",
                "Diligencia notas y asistencia en una planilla tipo celda. Administra el catálogo de cápsulas de su materia o logro.",
            ],
            [
                "Padre o acudiente",
                "Sigue una línea de tiempo con notas, asistencias y comunicados. En el centro de alertas confirma que se enteró y consulta la ruta de solución.",
            ],
            [
                "Estudiante",
                "Consulta horario, tareas y calificaciones. En el centro de nivelación completa las cápsulas asignadas y ve su barra de progreso.",
            ],
        ],
        [2800, 6800],
        aligns=["left", "left"],
    )
    add_body(
        doc,
        [
            (
                "El recorrido de una alerta tiene cuatro pasos. Una regla se cumple (por ejemplo, promedio de la materia menor a 3,0 en los últimos 15 días, o más de 3 fallas en el mes). El sistema avisa al acudiente y asigna la cápsula de esa competencia. El estudiante completa la guía, el video o el quiz. Al aprobar, el docente y el padre reciben el cambio de estado y la bitácora queda cerrada con la evidencia.",
                False,
            )
        ],
    )

    add_heading(doc, "8. ALCANCE DE ESTA VERSIÓN")
    add_table(
        doc,
        ["Módulo", "Qué incluye"],
        [
            [
                "1. Admisiones y matrículas",
                "Preinscripción en línea y carga de documentos. Legalización del cupo, grado y grupo. Ficha acumulativa. Expediente de acudientes, antecedentes médicos y contactos.",
            ],
            [
                "2. Control académico y asistencia",
                "Libro de calificaciones flexible según el SIEE. Asistencia diaria o por clase, con alerta al umbral definido. Boletines conceptuales y cuantitativos en PDF, masivos, con firma digital.",
            ],
            [
                "3. Alertas e intervención",
                "Reglas parametrizables de bajo rendimiento e inasistencia. Aviso por WhatsApp API, correo y notificación en el portal. Bitácora del Decreto 1290.",
            ],
            [
                "4. Refuerzo adaptativo",
                "Asignación de la cápsula de la competencia no superada. Guías, videos de YouTube, Vimeo o S3, y cuestionarios. Al aprobar, se actualiza el estado y se avisa al docente y al padre.",
            ],
            [
                "5. Rectoría",
                "Semáforo verde, amarillo y rojo por grados y materias. Roles y permisos de los seis perfiles. Reportes exportables.",
            ],
        ],
        [2800, 6800],
        aligns=["left", "left"],
    )
    add_subheading(doc, "Puesta en marcha")
    add_body(
        doc,
        [
            (
                "La configuración, la carga inicial acordada, la capacitación por rol y la documentación básica van incluidas en las tres modalidades. El calendario propuesto es de 10 a 14 semanas, contadas desde el día en que el colegio entrega el SIEE, la escala de valoración, los grados, los grupos y los datos de estudiantes y acudientes. Si esa información llega después, el calendario se corre.",
                False,
            )
        ],
    )
    add_subheading(doc, "Fuera de esta versión")
    add_body(
        doc,
        [
            (
                "Esta versión deja por fuera, y se cotiza cuando el colegio lo pida: cartera, pagos de pensión en línea, factura electrónica ante la DIAN, exportación o validación SIMAT, varias sedes en una sola cuenta, producción de las guías y los videos de las cápsulas, publicación en Play Store o App Store, y el consumo de conversaciones de WhatsApp. Ese consumo lo paga el colegio directamente a Meta. El contenido pedagógico de cada cápsula lo carga el colegio; Campus by TCD entrega el espacio, la asignación automática y el quiz.",
                False,
            )
        ],
    )

    add_heading(doc, "9. PRECIOS PUBLICADOS EN EL MERCADO")
    add_body(
        doc,
        [
            (
                "Las cifras de esta sección se leyeron el 13 de octubre de 2026 en las páginas y documentos que se citan. Pueden cambiar. Donde el origen no dice si el IVA va incluido, este informe tampoco lo afirma. Los ejemplos de 250 estudiantes marcados como cálculo de este informe son multiplicaciones propias, no cotizaciones de esos proveedores.",
                False,
            )
        ],
    )

    add_subheading(doc, "9.1. MiAula")
    add_body(
        doc,
        [
            (
                "Fuente: miaula.co/precio-software-gestion-escolar/. Tarifa base publicada, sin IVA y sin descuentos: $350 COP por estudiante al mes, o $3.150 COP por estudiante al año. La propia guía dice que el anual equivale a nueve meses de la tarifa mensual y mantiene el acceso doce meses. Incluye gestión académica, matrículas, usuarios y permisos, reportes, comunicación con familias, configuración inicial y formación por rol. Los mínimos se confirman con cada colegio.",
                False,
            )
        ],
    )
    add_table(
        doc,
        ["Estudiantes", "Mensual base", "Anual base"],
        [
            ["100", "$35.000", "$315.000"],
            ["250", "$87.500", "$787.500"],
            ["300", "$105.000", "$945.000"],
            ["500", "$175.000", "$1.575.000"],
        ],
        [3200, 3200, 3200],
        aligns=["center", "center", "center"],
    )
    add_note(
        doc,
        "MiAula publica los ejemplos de 100, 300 y 500 estudiantes y aclara que son aritmética, no cotizaciones. La fila de 250 es cálculo de este informe: 250 × $350 y 250 × $3.150.",
    )

    add_subheading(doc, "9.2. eScolary")
    add_body(
        doc,
        [
            (
                "Fuente: escolary.com.co/pricing. Planes desde 100 estudiantes activos. El valor sube o baja con la matrícula. En el ejemplo publicado de 250 estudiantes, Pro cuesta $2.000 por estudiante al mes ($500.000 al mes) y Plus cuesta $3.000 por estudiante al mes ($750.000 al mes). La página consultada no indica si esos valores incluyen IVA.",
                False,
            )
        ],
    )
    add_bullet(
        doc,
        [
            (
                "Pro: académico, matrículas, familias, pagos en línea con Wompi, cartera y exportación SIMAT.",
                False,
            )
        ],
    )
    add_bullet(
        doc,
        [
            (
                "Plus: todo Pro, avisos de WhatsApp de cartera, matrícula y calificaciones, varias sedes y soporte prioritario. eScolary indica que Plus paga a Meta directamente por las conversaciones de WhatsApp.",
                False,
            )
        ],
    )
    add_bullet(
        doc,
        [
            (
                "Por encima de 1.000 estudiantes, o para una red de colegios, eScolary pide una propuesta a la medida.",
                False,
            )
        ],
    )
    add_bullet(
        doc,
        [
            (
                "Plan Fundador para el año escolar 2027: implementación entre octubre y enero sin mensualidad, primer cobro en febrero de 2027, precio congelado 24 meses, migración incluida y plan anual en 10 cuotas. El plan mensual no exige permanencia.",
                False,
            )
        ],
    )

    add_subheading(doc, "9.3. Quid")
    add_body(
        doc,
        [
            (
                "Fuentes: quid.pw/planes-precio y el artículo de Equipo Quid del 6 de julio de 2026, actualizado el 16 de agosto de 2026. Quid publica precio por estudiante al año. Con pago anual anticipado se cobran 10 meses y se usan 12. Desde el plan Docente se suma un valor por docente. Las páginas consultadas no aclaran el IVA en el bloque de planes.",
                False,
            )
        ],
    )
    add_table(
        doc,
        ["Plan", "Por estudiante / año", "Por docente / año", "Mínimo anual"],
        [
            ["Esencial", "$2.000", "No aplica", "$400.000"],
            ["Docente", "$2.900", "$49.000", "$700.000"],
            ["Comunidad", "$4.900", "$49.000", "$1.500.000"],
            ["Institucional", "$9.900", "$99.000", "$3.000.000"],
        ],
        [2200, 2500, 2500, 2400],
        aligns=["left", "center", "center", "center"],
    )
    add_note(
        doc,
        "Cálculo de este informe para 250 estudiantes, antes de sumar docentes: Esencial 250 × $2.000 = $500.000 al año. Comunidad 250 × $4.900 = $1.225.000, por debajo del mínimo, así que aplica $1.500.000 al año. Institucional 250 × $9.900 = $2.475.000, por debajo del mínimo, así que aplica $3.000.000 al año. Los asientos docentes se suman a esas cifras.",
    )
    add_body(
        doc,
        [
            (
                "Esencial incluye núcleo académico, un administrador, acceso docente, portales de acudiente y estudiante, y generador de horarios. Comunidad suma anecdotario, elecciones y soporte humano por WhatsApp. Institucional incluye facturación electrónica, RFID, notificaciones por correo y soporte dedicado. El add-on de facturación electrónica cuesta $1.200 por estudiante al año, con mínimo de $390.000, y Quid aclara que la factura la emite el software contable del colegio, no Quid. Ese add-on figura gratis hasta el 31 de diciembre de 2026 y queda incluido en el plan Institucional.",
                False,
            )
        ],
    )

    add_subheading(doc, "9.4. Master2000")
    add_body(
        doc,
        [
            (
                "Master2000 no publica una lista por estudiante. Hay contratos y cotizaciones de colegios oficiales. La cotización 02760, del 23 de enero de 2026, para la IE Santa Catalina de Siena de Medellín, cobra $3.346.000 por la habilitación del sistema académico de un año lectivo 2026, con la aplicación incluida, excluido de IVA según el artículo 476 del Estatuto Tributario. En la misma cotización el asistente de bilingüismo vale $2.275.000 y el total aceptado el 27 de enero de 2026 es $5.621.000. Es el precio de un colegio y de un año, no una tarifa de catálogo.",
                False,
            )
        ],
    )
    add_note(
        doc,
        "Documento: modulo.master2000.net, cotización 02760, proceso de software académico 2026 de la IE Santa Catalina de Siena. Otro contrato público de 2026, el 002-2026 de la IE Antonio Donado Camacho con GMA Digital S.A.S., renueva la licencia académica en $2.095.000 por once meses. Se cita solo para mostrar que la licencia anual de ese proveedor se mueve por colegio, no por una tabla única.",
    )

    add_subheading(doc, "9.5. Proveedores que cotizan y no publican lista escolar")
    add_bullet(
        doc,
        [
            (
                "Phidias (phidias.com) reúne académico, financiero y comunicación con las familias, y cotiza la licencia según tamaño y módulos. En su página de servicios adicionales publica, entre otras tarifas, la capacitación extra a $60.000 por hora. Esa página no es el precio de la licencia.",
                False,
            )
        ],
    )
    add_bullet(
        doc,
        [
            (
                "Q10 publica precios para educación para el trabajo y el desarrollo humano. Esa tabla no corresponde a colegios de básica y media, y por eso este informe no la usa como tarifa escolar.",
                False,
            )
        ],
    )
    add_bullet(
        doc,
        [
            (
                "WebColegios y el plan completo de Sistema Saberes se cotizan. Quid lo registra así en sus comparativos públicos. Este informe no les asigna una cifra.",
                False,
            )
        ],
    )

    add_heading(doc, "10. LECTURA PARA CAMPUS BY TCD")
    add_body(
        doc,
        [
            (
                "Un colegio privado de 250 estudiantes, con las tarifas públicas de la sección 9, queda en órdenes de magnitud distintos según el producto que compre:",
                False,
            )
        ],
    )
    add_table(
        doc,
        ["Referencia", "Orden de valor en 250 estudiantes", "Qué está pagando"],
        [
            [
                "MiAula, tarifa base",
                "$87.500 al mes, o $787.500 al año",
                "Gestión académica, matrícula, reportes y comunicación con familias.",
            ],
            [
                "Quid Esencial",
                "$500.000 al año",
                "Núcleo académico y portales, con el mínimo de $400.000 superado.",
            ],
            [
                "Quid Comunidad",
                "$1.500.000 al año, más docentes",
                "Aplica el mínimo institucional. El cálculo por estudiante daba $1.225.000.",
            ],
            [
                "eScolary Pro",
                "$500.000 al mes",
                "Académico, familias, cartera, pagos en línea y SIMAT.",
            ],
            [
                "eScolary Plus",
                "$750.000 al mes",
                "Pro, más WhatsApp, varias sedes y soporte prioritario.",
            ],
            [
                "Master2000, cotización 02760",
                "$3.346.000 por un año académico",
                "Un colegio oficial, un año. No es un ejemplo de 250 estudiantes.",
            ],
        ],
        [2300, 3400, 3900],
        aligns=["left", "left", "left"],
    )
    add_body(
        doc,
        [
            (
                "La recomendación de la sección 11 deja el plan Aula de Campus by TCD en $750.000 al mes para un colegio de hasta 300 estudiantes. En el ejemplo público de eScolary, ese es el mismo valor mensual de Plus para 250 estudiantes. MiAula y Quid publican la gestión académica muy por debajo, porque su tarifa visible cubre matrícula, notas y portales. Campus by TCD agrega el motor de reglas, la bitácora del Decreto 1290 y el refuerzo con cápsulas.",
                False,
            )
        ],
    )
    add_body(
        doc,
        [
            (
                "eScolary Plus, en ese mismo valor, ya incluye cartera, pagos, SIMAT y varias sedes. Esos procesos quedan por fuera de esta versión de Campus by TCD. La directiva aprueba, con la tabla de la sección 11, un precio de intervención y permanencia, y deja la tesorería para una cotización posterior.",
                False,
            )
        ],
    )
    add_body(
        doc,
        [
            (
                "Hay un salto entre bandas: un colegio de 300 estudiantes permanece en Aula ($750.000 al mes) y uno de 301 entra a Colegio ($1.320.000 al mes). La sección 16 pide a la directiva confirmar ese salto o pedir, más adelante, una regla por estudiante dentro de la banda.",
                False,
            )
        ],
    )

    add_heading(doc, "11. MODALIDADES Y PRECIOS PROPUESTOS")
    add_body(
        doc,
        [
            (
                "Los valores de esta sección son la recomendación de este informe para aprobación de la directiva. No son precios publicados del mercado. Van en pesos colombianos. El IVA se liquida según el régimen de ",
                False,
            ),
            ("THE COMPANY DEV", True),
            (" al momento de facturar, y no está sumado en las tablas.", False),
        ],
    )
    add_body(
        doc,
        [
            (
                "Campus by TCD se ofrece en tres modalidades. En ninguna se cobra una implementación aparte: la configuración, la puesta en marcha y la capacitación por rol están incluidas.",
                False,
            )
        ],
    )
    add_table(
        doc,
        ["Concepto", "Mensual", "Anualidad", "Licencia vitalicia"],
        [
            ["Forma de pago", "Cada mes", "Un pago al año", "Pago único, 50 % y 50 %"],
            ["Implementación", "Incluida", "Incluida", "Incluida"],
            ["Uso del sistema", "Mientras pague", "12 meses", "Indefinido, versión entregada"],
            ["Soporte de 3 h/mes", "Incluido", "Incluido", "3 meses; luego plan opcional"],
            ["Actualizaciones", "Incluidas", "Incluidas", "Con el plan de soporte"],
            ["WhatsApp API", "Consumo del colegio ante Meta", "Consumo del colegio ante Meta", "Consumo del colegio ante Meta"],
            ["Hosting, dominio o VPS", "Del cliente", "Del cliente", "Del cliente"],
            ["Código fuente", "THE COMPANY DEV", "THE COMPANY DEV", "THE COMPANY DEV"],
            ["Beneficio", "1 mes gratis por año cumplido", "Se pagan 10 meses y se usan 12", "Sin canon de licencia"],
        ],
        [2100, 2500, 2500, 2500],
        aligns=["left", "center", "center", "center"],
        size=9,
        header_size=9,
    )

    add_subheading(doc, "11.1. Suscripción mensual")
    add_body(
        doc,
        [
            (
                "Cada plan es un valor fijo hasta un tope de estudiantes activos. Si la matrícula supera el tope, el colegio pasa al plan siguiente.",
                False,
            )
        ],
    )
    add_table(
        doc,
        ["Plan", "Estudiantes activos", "Mensualidad", "Valor por estudiante en el tope"],
        [
            ["Semilla", "Hasta 150", "$450.000", "$3.000"],
            ["Aula", "Hasta 300", "$750.000", "$2.500"],
            ["Colegio", "Hasta 600", "$1.320.000", "$2.200"],
            ["Red", "Más de 600 o varias sedes", "Desde $1.800.000", "Se confirma en la propuesta"],
        ],
        [1600, 2800, 2400, 2800],
        aligns=["left", "center", "center", "center"],
    )
    add_note(
        doc,
        "El valor por estudiante en el tope es cálculo de este informe: la mensualidad dividida por el máximo del plan. Por debajo del tope, el valor por estudiante sube, porque el plan es un fijo.",
    )

    add_subheading(doc, "11.2. Anualidad (se pagan 10 meses y se usan 12)")
    add_table(
        doc,
        ["Plan", "Pago anual", "Equivale a", "Ahorro frente a 12 meses"],
        [
            ["Semilla", "$4.500.000", "$375.000 / mes", "$900.000"],
            ["Aula", "$7.500.000", "$625.000 / mes", "$1.500.000"],
            ["Colegio", "$13.200.000", "$1.100.000 / mes", "$2.640.000"],
            ["Red", "Desde $18.000.000", "$1.500.000 / mes", "$3.600.000 sobre la base"],
        ],
        [1600, 2500, 2500, 3000],
        aligns=["left", "center", "center", "center"],
    )
    add_note(
        doc,
        "Cálculo de este informe: pago anual = mensualidad × 10. Equivale a = pago anual ÷ 12. Ahorro = mensualidad × 2. En Red, el ahorro corresponde a la base de $1.800.000.",
    )

    add_subheading(doc, "11.3. Licencia vitalicia (pago único)")
    add_body(
        doc,
        [
            (
                "El pago único propuesto equivale a 22 mensualidades del plan, con redondeo en Colegio y en Red. Cubre el derecho de uso de la versión entregada.",
                False,
            )
        ],
    )
    add_table(
        doc,
        ["Plan", "Pago único", "Plan de soporte opcional"],
        [
            ["Semilla", "$9.900.000", "$220.000 / mes"],
            ["Aula", "$16.500.000", "$320.000 / mes"],
            ["Colegio", "$29.000.000", "$480.000 / mes"],
            ["Red", "Desde $40.000.000", "Se cotiza"],
        ],
        [2400, 3600, 3600],
        aligns=["left", "center", "center"],
    )
    add_bullet(
        doc,
        [
            (
                "Incluye el derecho de uso indefinido de la versión entregada, la instalación en el hosting o VPS del cliente, la capacitación y 3 meses de soporte de garantía.",
                False,
            )
        ],
    )
    add_bullet(
        doc,
        [
            (
                "Después de la garantía, el cliente asume su propio soporte. Si toma el plan opcional, recibe 3 horas al mes y las actualizaciones.",
                False,
            )
        ],
    )
    add_bullet(
        doc,
        [
            (
                "El código fuente sigue siendo de ",
                False,
            ),
            ("THE COMPANY DEV", True),
            (
                ". La licencia da derecho de uso. No autoriza reventa, copia ni modificación por terceros.",
                False,
            ),
        ],
    )

    add_subheading(doc, "11.4. Servicios que se cotizan aparte")
    add_table(
        doc,
        ["Servicio", "Valor propuesto"],
        [
            [
                "Hora adicional de soporte o ajustes",
                "$90.000. Es la misma hora ya planteada en el perfil de RecoRoute.",
            ],
            [
                "Cápsulas producidas por THE COMPANY DEV",
                "Se cotizan. El precio del plan cubre la plataforma, no la producción del contenido.",
            ],
            ["Cartera, pagos en línea, SIMAT o factura DIAN", "Se cotizan según el alcance."],
            ["Sede adicional dentro de un plan que no sea Red", "Se cotiza, o el colegio pasa al plan Red."],
            ["Publicación en Play Store o App Store", "Se cotiza según el alcance."],
            ["Consumo de WhatsApp API", "Lo paga el colegio a Meta. No va en la mensualidad."],
        ],
        [4200, 5400],
        aligns=["left", "left"],
    )

    add_heading(doc, "12. COMPARACIÓN DE ALCANCE FRENTE AL MERCADO")
    add_table(
        doc,
        ["Proceso", "Campus by TCD, esta versión", "Dónde aparece en el mercado citado"],
        [
            ["Matrícula y expediente", "Incluido", "MiAula, eScolary, Quid y Phidias."],
            ["Calificaciones, SIEE y boletines", "Incluido", "MiAula, eScolary, Quid y Master2000."],
            ["Asistencia con umbral de alerta", "Incluido", "La alerta automática de fallas es parte del diferencial de Campus."],
            ["Reglas de bajo rendimiento", "Incluidas y parametrizables", "No forman parte de las tarifas públicas citadas."],
            ["Bitácora del Decreto 1290", "Incluida", "No forma parte de las tarifas públicas citadas."],
            ["Cápsulas y quiz de nivelación", "Incluidos", "No forman parte de las tarifas públicas citadas."],
            ["Semáforo de rectoría", "Incluido", "Los reportes existen en el mercado. El semáforo de riesgo es de Campus."],
            ["Cartera y pagos en línea", "Fuera de esta versión", "eScolary Pro y Plus, Phidias y el módulo financiero de Quid."],
            ["SIMAT", "Fuera de esta versión", "eScolary Pro y Plus."],
            ["Factura electrónica DIAN", "Fuera de esta versión", "Quid, por medio del software contable. Phidias la anuncia."],
            ["Consumo de WhatsApp", "Lo paga el colegio a Meta", "eScolary Plus también lo deja en cabeza del colegio."],
        ],
        [2500, 3300, 3800],
        aligns=["left", "left", "left"],
        size=9,
        header_size=9,
    )

    add_heading(doc, "13. EJEMPLO DE INVERSIÓN")
    add_body(
        doc,
        [
            (
                "Colegio privado de 400 estudiantes activos. Entra al plan Colegio, porque supera el tope de 300 de Aula y queda dentro del tope de 600. Sin servicios adicionales. Cálculo de este informe.",
                False,
            )
        ],
    )
    add_table(
        doc,
        ["Modalidad", "Primer pago", "Pagos siguientes", "Total del primer año"],
        [
            ["Mensual", "$1.320.000", "$1.320.000 cada mes", "$15.840.000"],
            ["Anualidad", "$13.200.000", "$13.200.000 al renovar", "$13.200.000"],
            [
                "Licencia vitalicia",
                "$14.500.000 (50 %)",
                "$14.500.000 a la entrega",
                "$29.000.000",
            ],
            [
                "Vitalicia con soporte",
                "$14.500.000 (50 %)",
                "Segundo 50 % a la entrega y $480.000 al mes desde el mes 4",
                "$33.320.000",
            ],
        ],
        [2000, 2300, 2900, 2400],
        aligns=["left", "center", "left", "center"],
        size=9,
        header_size=9,
    )
    add_note(
        doc,
        "La entrega queda dentro de las 10 a 14 semanas, así que los dos pagos de la licencia vitalicia caen en el primer año: $29.000.000. El soporte opcional del ejemplo son nueve meses, de abril a diciembre si el arranque es en enero: 9 × $480.000 = $4.320.000. Sumados a la licencia dan $33.320.000.",
    )
    add_body(
        doc,
        [
            (
                "En el tope del plan Colegio, $1.320.000 repartidos en 600 estudiantes equivalen a $2.200 por estudiante al mes. En este ejemplo de 400 estudiantes, el fijo equivale a $3.300 por estudiante al mes (cálculo de este informe: $1.320.000 ÷ 400). Sigue por debajo de los $3.000 de eScolary Plus solo cuando el colegio se acerca al tope de 600; en 400 estudiantes el fijo queda por encima de esa referencia por estudiante, y muy por encima de MiAula.",
                False,
            )
        ],
    )

    add_heading(doc, "14. CONDICIONES DEL SERVICIO")
    add_table(
        doc,
        ["Condición", "Detalle"],
        [
            ["Horario de soporte", "Lunes a viernes, de 9:00 a.m. a 5:00 p.m., por correo y WhatsApp."],
            ["Tiempo de respuesta", "Primera respuesta en máximo 24 horas hábiles."],
            ["Horas incluidas", "3 horas al mes, no acumulables. Aplican a la mensualidad, a la anualidad y al plan de soporte."],
            ["Pago mensual", "Dentro de los primeros 5 días de cada mes."],
            ["Pago anual", "Un solo pago al iniciar el periodo."],
            ["Pago de la licencia vitalicia", "50 % al iniciar y 50 % a la entrega."],
            ["Medios de pago", "Nequi, transferencia o efectivo."],
            ["Fidelización", "1 mes gratis por cada año completo en la suscripción mensual, entregado al final del año. No aplica a la anualidad ni a la licencia vitalicia."],
            ["Crecimiento", "Si la matrícula supera el plan, el colegio pasa al plan siguiente. Un módulo nuevo se cotiza cuando lo pidan."],
            ["Información del colegio", "El calendario de 10 a 14 semanas corre cuando llegan el SIEE, los grados, los grupos y los datos."],
        ],
        [2800, 6800],
        aligns=["left", "left"],
    )

    add_heading(doc, "15. RESPONSABILIDADES Y PROPIEDAD INTELECTUAL")
    add_bullet(
        doc,
        [
            ("Hosting, dominio o VPS: ", False),
            ("EL CLIENTE", True),
            (" los contrata, los paga y los mantiene. ", False),
            ("THE COMPANY DEV", True),
            (
                " no responde por caídas, lentitud, pérdida de información, vencimientos, renovaciones ni fallas del proveedor. Sí entrega los requisitos mínimos del servidor.",
                False,
            ),
        ],
    )
    add_bullet(
        doc,
        [
            ("Certificado SSL y correos institucionales: a cargo de ", False),
            ("EL CLIENTE", True),
            (".", False),
        ],
    )
    add_bullet(
        doc,
        [
            ("Información académica: ", False),
            ("EL CLIENTE", True),
            (
                " entrega a tiempo el SIEE, la escala, los grados, los grupos, los estudiantes, los acudientes y los accesos. También carga el contenido de las cápsulas. Los retrasos mueven el calendario.",
                False,
            ),
        ],
    )
    add_bullet(
        doc,
        [
            ("Línea de WhatsApp y consumo ante Meta: a cargo de ", False),
            ("EL CLIENTE", True),
            (". ", False),
            ("THE COMPANY DEV", True),
            (" conecta el envío cuando el colegio tiene la línea y la API habilitadas.", False),
        ],
    )
    add_bullet(
        doc,
        [
            ("Código fuente: es propiedad de ", False),
            ("THE COMPANY DEV", True),
            (" en todas las modalidades.", False),
        ],
    )
    add_bullet(
        doc,
        [
            (
                "Datos personales: se tratan conforme a la Ley 1581 de 2012. La política de tratamiento y el aviso de privacidad del colegio son responsabilidad de ",
                False,
            ),
            ("EL CLIENTE", True),
            (".", False),
        ],
    )

    add_heading(doc, "16. PUNTOS PARA APROBACIÓN DE LA DIRECTIVA")
    add_bullet(
        doc,
        [("Aprobar los cuatro planes y las tres modalidades de la sección 11, con la implementación incluida.", False)],
    )
    add_bullet(
        doc,
        [
            (
                "Confirmar que el plan Aula queda en el orden de $750.000 al mes, alineado con el ejemplo público de eScolary Plus para 250 estudiantes, sabiendo que cartera, SIMAT y factura electrónica quedan por fuera.",
                False,
            )
        ],
    )
    add_bullet(
        doc,
        [
            (
                "Confirmar el salto de Aula a Colegio al pasar de 300 a 301 estudiantes, o pedir una regla por estudiante dentro de la banda antes de usar la tabla con un cliente.",
                False,
            )
        ],
    )
    add_bullet(
        doc,
        [
            (
                "Definir si los tres primeros colegios congelan por 24 meses el valor del plan que firmen, sin bajar la tarifa, tomando como referencia comercial el plan fundador que eScolary publicó para 2027.",
                False,
            )
        ],
    )
    add_bullet(
        doc,
        [("Aprobar el calendario de 10 a 14 semanas y las 3 horas mensuales de soporte.", False)],
    )
    add_bullet(
        doc,
        [("Aprobar la hora adicional de $90.000, en la misma referencia del perfil de RecoRoute.", False)],
    )
    add_bullet(
        doc,
        [
            (
                "Dejar por escrito, en cada propuesta a un colegio, que el consumo de WhatsApp lo paga el colegio a Meta y que el contenido de las cápsulas lo carga el colegio.",
                False,
            )
        ],
    )
    add_body(
        doc,
        [("Elaborado por", True)],
        before=14,
        after=0,
        align="left",
        keep_next=True,
    )
    add_body(
        doc,
        [("Jhon Javier Angulo Hilarion", False)],
        before=0,
        after=0,
        align="left",
        keep_next=True,
    )
    add_body(
        doc,
        [("Ejecutivo Comercial", False)],
        before=0,
        after=0,
        align="left",
        keep_next=True,
    )
    add_body(
        doc,
        [("THE COMPANY DEV", True)],
        before=0,
        after=0,
        align="left",
        keep_next=True,
    )
    add_contacts(doc)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
